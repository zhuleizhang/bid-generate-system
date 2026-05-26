"""招标文件解析引擎 — PDF（pdfplumber + pypdf 降级）和 Word（python-docx）。"""

import io
import json
import re
from typing import Any

import pdfplumber
from docx import Document as DocxDocument
from pypdf import PdfReader

# 中文标书常见的章节标题模式
_HEADING_PATTERNS = [
    (re.compile(r"^第[一二三四五六七八九十百千\d]+[章节条]"), 1),
    (re.compile(r"^[一二三四五六七八九十]+[、，]"), 1),
    (re.compile(r"^（[一二三四五六七八九十]+）"), 2),
    (re.compile(r"^\([一二三四五六七八九十]+\)"), 2),
    (re.compile(r"^\d+\.\d+\.\d+"), 4),
    (re.compile(r"^\d+\.\d+"), 3),
    (re.compile(r"^\d+[、，.)]"), 2),
    (re.compile(r"^[一二三四五六七八九十]+[）\)]"), 1),
]


def _detect_heading_level(line: str) -> int:
    """根据行首模式判断标题层级，非标题返回 0。"""
    for pattern, level in _HEADING_PATTERNS:
        if pattern.match(line):
            return level
    return 0


def _detect_sections_from_text(text: str) -> list[dict[str, Any]]:
    """从纯文本中检测章节结构。"""
    sections: list[dict[str, Any]] = []
    lines = text.split("\n")
    current_section: dict[str, Any] | None = None

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        level = _detect_heading_level(stripped)
        # 短行 + 匹配标题模式视为章节标题
        is_heading = level > 0 and len(stripped) < 120

        if is_heading:
            if current_section and (current_section["paragraphs"] or current_section["title"] != "文档开头"):
                sections.append(current_section)
            current_section = {
                "title": stripped,
                "level": level,
                "content": "",
                "paragraphs": [],
            }
        elif current_section is not None:
            current_section["paragraphs"].append(stripped)
            current_section["content"] += stripped + "\n"
        else:
            current_section = {
                "title": "文档开头",
                "level": 0,
                "content": stripped + "\n",
                "paragraphs": [stripped],
            }

    if current_section and (current_section["paragraphs"] or current_section.get("content")):
        sections.append(current_section)

    return sections


def parse_pdf(file_bytes: bytes) -> dict[str, Any]:
    """解析 PDF 文件，pdfplumber 为主，pypdf 为降级方案。"""
    try:
        return _parse_with_pdfplumber(file_bytes)
    except Exception:
        return _parse_with_pypdf(file_bytes)


def _parse_with_pdfplumber(file_bytes: bytes) -> dict[str, Any]:
    """使用 pdfplumber 提取文本和表格。"""
    sections_list: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    all_text: list[str] = []

    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page_num, page in enumerate(pdf.pages, 1):
            text = page.extract_text()
            if text:
                all_text.append(text)

            page_tables = page.extract_tables()
            for t_idx, table in enumerate(page_tables):
                if not table or not table[0]:
                    continue
                headers = [str(c) if c else "" for c in table[0]]
                rows = [[str(c) if c else "" for c in row] for row in table[1:]]
                tables.append({
                    "caption": f"第{page_num}页 表格{len(tables) + 1}",
                    "headers": headers,
                    "rows": rows,
                })

    full_text = "\n".join(all_text)
    sections_list = _detect_sections_from_text(full_text)

    return {
        "sections": sections_list,
        "tables": tables,
        "full_text": full_text,
    }


def _parse_with_pypdf(file_bytes: bytes) -> dict[str, Any]:
    """使用 pypdf 作为降级方案，仅提取纯文本。"""
    all_text: list[str] = []
    reader = PdfReader(io.BytesIO(file_bytes))

    for page in reader.pages:
        text = page.extract_text()
        if text:
            all_text.append(text)

    full_text = "\n".join(all_text)
    sections_list = _detect_sections_from_text(full_text)

    return {
        "sections": sections_list,
        "tables": [],
        "full_text": full_text,
    }


def parse_docx_text(file_bytes: bytes) -> dict[str, Any]:
    """使用 python-docx 解析 Word 文件，提取段落和表格文本。"""
    doc = DocxDocument(io.BytesIO(file_bytes))

    sections_list: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    all_text: list[str] = []

    current_section: dict[str, Any] | None = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        all_text.append(text)
        style_name = para.style.name if para.style else ""

        if style_name.startswith("Heading"):
            if current_section is not None:
                sections_list.append(current_section)
            # Heading 1 → level 1, Heading 2 → level 2, ...
            level = 1
            for char in style_name:
                if char.isdigit():
                    level = int(char)
                    break
            current_section = {
                "title": text,
                "level": level,
                "content": "",
                "paragraphs": [],
            }
        elif current_section is not None:
            current_section["paragraphs"].append(text)
            current_section["content"] += text + "\n"
        else:
            current_section = {
                "title": "文档开头",
                "level": 0,
                "content": text + "\n",
                "paragraphs": [text],
            }

    if current_section is not None:
        sections_list.append(current_section)

    # 如果没有检测到任何 Heading 样式，回退为文本模式检测
    has_headings = any(s["level"] > 0 for s in sections_list)
    if not has_headings:
        full_text = "\n".join(all_text)
        sections_list = _detect_sections_from_text(full_text)

    for idx, table in enumerate(doc.tables):
        headers: list[str] = []
        rows: list[list[str]] = []
        for row_idx, row in enumerate(table.rows):
            cells = [cell.text for cell in row.cells]
            if row_idx == 0:
                headers = cells
            else:
                rows.append(cells)
        tables.append({
            "caption": f"表格{idx + 1}",
            "headers": headers,
            "rows": rows,
        })

    return {
        "sections": sections_list,
        "tables": tables,
        "full_text": "\n".join(all_text),
    }


async def parse_tender_document(doc_id: str) -> dict[str, Any]:
    """解析单个招标文件并存储结果到 document_versions 表。

    根据文档的 mime_type 选择 PDF 或 Word 解析引擎，
    解析结果存储为结构化 JSON（sections + tables + full_text），
    同时写入 document_versions（version_type = 'parsed'）。
    """
    from app.gateway.supabase_gateway import SupabaseGateway

    gw = SupabaseGateway()
    doc = gw.get_document(doc_id)
    if not doc:
        raise ValueError(f"文档不存在: {doc_id}")

    file_path = doc.get("file_path", "")
    if not file_path:
        raise ValueError("文档缺少 file_path")

    mime_type = doc.get("mime_type", "")
    doc_name = doc.get("name", "")

    # 下载文件内容
    file_bytes = gw.download_file_by_url(file_path)

    # 根据 mime_type 选择解析引擎
    if mime_type == "application/pdf":
        result = parse_pdf(file_bytes)
        file_type = "pdf"
    elif mime_type in (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ):
        result = parse_docx_text(file_bytes)
        file_type = "docx"
    else:
        raise ValueError(f"不支持的文件类型: {mime_type}，仅支持 PDF 和 Word")

    # 将结构化结果存储到 document_versions
    content_json = json.dumps(result, ensure_ascii=False)
    metadata = {
        "file_type": file_type,
        "section_count": len(result["sections"]),
        "table_count": len(result["tables"]),
        "text_length": len(result["full_text"]),
    }

    version = gw.insert_document_version({
        "document_id": doc_id,
        "version_type": "parsed",
        "content": content_json,
        "metadata": metadata,
    })

    # 更新文档状态
    gw.update_document_status(doc_id, "parsed")

    return {
        "document_id": doc_id,
        "document_name": doc_name,
        "file_type": file_type,
        "status": "success",
        "version_id": version.get("id", ""),
        "sections": result["sections"],
        "tables": result["tables"],
        "full_text": result["full_text"],
    }
